#!/usr/bin/env python3
"""
Vani Portal Data Indexer
Extracts and normalizes all railway track assets, personnel, store inventory,
engineering drawings, and longitudinal elevation/gradient profiles into a single unified JSON database.
"""

import os
import re
import glob
import json
from pathlib import Path
import openpyxl
import xlrd

BASE_DIR = Path('/Users/vivekazad/Desktop/Vani')
DATA_DIR = BASE_DIR / 'data'
OUTPUT_FILE = DATA_DIR / 'vani_database.json'

def clean_val(val):
    if val is None:
        return ""
    if isinstance(val, float):
        if val.is_integer():
            return int(val)
        return round(val, 4)
    s = str(val).strip()
    if s == "None" or s == "#VALUE!":
        return ""
    return s

def safe_float(val, default=0.0):
    if val is None:
        return default
    try:
        s = str(val).strip().replace(',', '')
        m = re.search(r'[-+]?\d*\.?\d+', s)
        if m:
            return float(m.group(0))
        return default
    except Exception:
        return default

def parse_at_a_glance():
    xlsx_path = BASE_DIR / 'At a Glance.xlsx'
    if not xlsx_path.exists():
        xlsx_path = DATA_DIR / 'Store' / 'At a Glance.xlsx'
    
    print(f"Loading {xlsx_path}...")
    wb = openpyxl.load_workbook(xlsx_path, data_only=True, read_only=True)
    
    results = {}

    # 1. Section Summary (At a Glance)
    if 'At a Glance' in wb.sheetnames:
        ws = wb['At a Glance']
        sections = []
        for row in ws.iter_rows(values_only=True):
            cells = [clean_val(c) for c in row]
            if len(cells) >= 11 and cells[0] and cells[0] not in ['Section', 'Section under IMSD SMUN (Bridge Details)']:
                sections.append({
                    'section': str(cells[0]),
                    'line': str(cells[1]),
                    'km_from': safe_float(cells[2]),
                    'km_to': safe_float(cells[3]),
                    'mjb': int(safe_float(cells[4])),
                    'mib': int(safe_float(cells[5])),
                    'rub': int(safe_float(cells[6])),
                    'rob': int(safe_float(cells[7])),
                    'fob': int(safe_float(cells[8])),
                    'owg': int(safe_float(cells[9])),
                    'lc': int(safe_float(cells[10])) if len(cells) > 10 else 0
                })
        results['sections'] = sections

    # 2. Bridges
    if 'Bridge' in wb.sheetnames:
        ws = wb['Bridge']
        bridges = []
        rows = list(ws.iter_rows(values_only=True))
        for r in rows[1:]:
            cells = [clean_val(c) for c in r]
            if any(cells) and len(cells) > 1 and cells[1]:
                bridges.append({
                    's_no': cells[0],
                    'bridge_no': str(cells[1]),
                    'section': str(cells[2]) if len(cells) > 2 else '',
                    'km_from': safe_float(cells[3]) if len(cells) > 3 else 0.0,
                    'km_to': safe_float(cells[4]) if len(cells) > 4 else 0.0,
                    'bridge_type': str(cells[5]) if len(cells) > 5 else '',
                    'old_bridge_no': str(cells[6]) if len(cells) > 6 else '',
                    'category': str(cells[7]) if len(cells) > 7 else '',
                    'span_config': str(cells[8]) if len(cells) > 8 else '',
                    'length': safe_float(cells[9]) if len(cells) > 9 else 0.0,
                    'linear_waterway': str(cells[10]) if len(cells) > 10 else '',
                    'remarks': str(cells[11]) if len(cells) > 11 else ''
                })
        results['bridges'] = bridges

    # 3. Curves
    if 'Curve' in wb.sheetnames:
        ws = wb['Curve']
        curves = []
        rows = list(ws.iter_rows(values_only=True))
        for r in rows[2:]:
            cells = [clean_val(c) for c in r]
            if len(cells) >= 10 and cells[1]:
                curves.append({
                    's_no': cells[0],
                    'curve_no': str(cells[1]),
                    'km_from': safe_float(cells[2]),
                    'km_to': safe_float(cells[3]),
                    'length': safe_float(cells[4]),
                    'degree': safe_float(cells[5]),
                    'radius_tms': safe_float(cells[6]),
                    'radius': safe_float(cells[7]),
                    'speed': safe_float(cells[8]),
                    'transition_length': safe_float(cells[9]),
                    'circular_length': safe_float(cells[10]) if len(cells) > 10 else 0,
                    'cant_se': safe_float(cells[12]) if len(cells) > 12 else 0,
                    'remarks': str(cells[13]) if len(cells) > 13 else ''
                })
        results['curves'] = curves

    # 4. Points and Crossings (PNC)
    if 'PNC' in wb.sheetnames:
        ws = wb['PNC']
        pnc = []
        rows = list(ws.iter_rows(values_only=True))
        for r in rows[2:]:
            cells = [clean_val(c) for c in r]
            if len(cells) >= 6 and cells[2]:
                pnc.append({
                    's_no': cells[0],
                    'station': str(cells[1]),
                    'point_no': str(cells[2]),
                    'line': str(cells[3]),
                    'angle': str(cells[4]),
                    'srj_chainage': safe_float(cells[5]),
                    'laid_on': str(cells[6]) if len(cells) > 6 else 'Straight',
                    'turnout': str(cells[7]) if len(cells) > 7 else '',
                    'traffic': str(cells[8]) if len(cells) > 8 else '',
                    'se': str(cells[9]) if len(cells) > 9 else '',
                    'd': str(cells[10]) if len(cells) > 10 else '',
                    'stations_behind': str(cells[11]) if len(cells) > 11 else ''
                })
        results['points_and_crossings'] = pnc

    # 5. LC, LWR, SEJ
    if 'LC, LWR, SEJ' in wb.sheetnames:
        ws = wb['LC, LWR, SEJ']
        lwr_list = []
        rows = list(ws.iter_rows(values_only=True))
        for r in rows[2:]:
            cells = [clean_val(c) for c in r]
            if len(cells) >= 6 and cells[1]:
                lwr_list.append({
                    's_no': cells[0],
                    'lwr_no': str(cells[1]),
                    'section': str(cells[2]),
                    'km_from': safe_float(cells[3]),
                    'km_to': safe_float(cells[4]),
                    'length': safe_float(cells[5]),
                    'gap_measured_on': str(cells[6]) if len(cells) > 6 else ''
                })
        results['lwr_sej'] = lwr_list

    # 6. DFWO (Defective Welds / Rails)
    if 'DFWO' in wb.sheetnames:
        ws = wb['DFWO']
        dfwo_list = []
        rows = list(ws.iter_rows(values_only=True))
        for r in rows[3:]:
            cells = [clean_val(c) for c in r]
            if len(cells) >= 6 and (cells[1] or cells[3]):
                km = safe_float(cells[1])
                m = safe_float(cells[2])
                ch = km + (m / 1000.0) if km else 0.0
                dfwo_list.append({
                    's_no': cells[0],
                    'km': km,
                    'meter': m,
                    'chainage': round(ch, 4),
                    'defect_no': str(cells[3]),
                    'line_type': str(cells[4]),
                    'remark': str(cells[5])
                })
        results['dfwo'] = dfwo_list

    # 7. Turnout Defects
    if 'Defects' in wb.sheetnames:
        ws = wb['Defects']
        defects = []
        rows = list(ws.iter_rows(values_only=True))
        for r in rows[2:]:
            cells = [clean_val(c) for c in r]
            if len(cells) >= 6 and cells[2]:
                defects.append({
                    's_no': cells[0],
                    'station': str(cells[1]),
                    'point_no': str(cells[2]),
                    'line': str(cells[3]),
                    'angle': str(cells[4]),
                    'srj_chainage': safe_float(cells[5]),
                    'wear_nose_lh': safe_float(cells[6]) if len(cells) > 6 else 0,
                    'wear_nose_center': safe_float(cells[7]) if len(cells) > 7 else 0,
                    'wear_nose_rh': safe_float(cells[8]) if len(cells) > 8 else 0,
                    'wear_269_lh': safe_float(cells[9]) if len(cells) > 9 else 0,
                    'wear_269_center': safe_float(cells[10]) if len(cells) > 10 else 0,
                    'wear_269_rh': safe_float(cells[11]) if len(cells) > 11 else 0,
                    'remark': str(cells[12]) if len(cells) > 12 else ''
                })
        results['turnout_defects'] = defects

    # 8. Permanent / Officers Staff
    if 'Permanent Staff' in wb.sheetnames:
        ws = wb['Permanent Staff']
        perm_staff = []
        rows = list(ws.iter_rows(values_only=True))
        for r in rows[4:]:
            cells = [clean_val(c) for c in r]
            if len(cells) >= 5 and cells[1]:
                perm_staff.append({
                    'name': str(cells[1]),
                    'designation': str(cells[2]),
                    'emp_id': str(cells[3]),
                    'posting': str(cells[4]),
                    'mobile': str(cells[5]) if len(cells) > 5 else '',
                    'email': str(cells[6]) if len(cells) > 6 else ''
                })
        results['permanent_staff'] = perm_staff

    # 9. Master Staff (ExMaster_DATA)
    if 'ExMaster_DATA' in wb.sheetnames:
        ws = wb['ExMaster_DATA']
        master_staff = []
        rows = list(ws.iter_rows(values_only=True))

        def parse_staff_chainage(text):
            if not text:
                return 0.0, 0.0
            s = str(text).strip()
            s = re.sub(r'1078\b', '1178', s)  # Correct obvious typo in sheet
            if '-' in s or 'to' in s.lower():
                parts = re.split(r'[-–]|(?:\bto\b)', s, flags=re.IGNORECASE)
                if len(parts) == 2:
                    def parse_p(p):
                        p = p.strip()
                        slash = re.search(r'(\d+)/(\d+)', p)
                        if slash:
                            km = float(slash.group(1))
                            tp = float(slash.group(2))
                            return km + (tp / 20.0 if tp <= 20 else tp / 1000.0)
                        dec = re.search(r'(\d+\.?\d*)', p)
                        if dec:
                            return float(dec.group(1))
                        return 0.0
                    k1 = parse_p(parts[0])
                    k2 = parse_p(parts[1])
                    return round(min(k1, k2), 3), round(max(k1, k2), 3)
            km_m = re.findall(r'(\d+\.?\d*)', s)
            k1 = float(km_m[0]) if len(km_m) >= 1 else 0.0
            k2 = float(km_m[1]) if len(km_m) >= 2 else 0.0
            return round(min(k1, k2), 3), round(max(k1, k2), 3)

        for r in rows[2:]:
            cells = [clean_val(c) for c in r]
            if len(cells) >= 8 and (cells[1] or cells[2]):
                awpo_id = str(cells[1])
                name = str(cells[2])
                km_text = str(cells[5])
                km_from, km_to = parse_staff_chainage(km_text)

                master_staff.append({
                    's_no': cells[0],
                    'awpo_id': awpo_id,
                    'name': name,
                    'father_name': str(cells[3]),
                    'beat_no': str(cells[4]),
                    'km_range': km_text,
                    'km_from': km_from,
                    'km_to': km_to,
                    'mobile': str(cells[6]),
                    'designation': str(cells[7]),
                    'residence': str(cells[8]),
                    'district': str(cells[9]),
                    'other_contact': str(cells[10]) if len(cells) > 10 else '',
                    'email': str(cells[13]) if len(cells) > 13 else ''
                })
        results['master_staff'] = master_staff

    # 10. Gatemen
    if 'Gateman' in wb.sheetnames:
        ws = wb['Gateman']
        gatemen = []
        rows = list(ws.iter_rows(values_only=True))
        for r in rows[2:]:
            cells = [clean_val(c) for c in r]
            if len(cells) >= 6 and cells[1]:
                gatemen.append({
                    's_no': cells[0],
                    'name': str(cells[1]),
                    'id_no': str(cells[2]),
                    'lc_no': str(cells[3]),
                    'section': str(cells[4]),
                    'mobile': str(cells[5]),
                    'rest_giver': str(cells[6]) if len(cells) > 6 else ''
                })
        results['gatemen'] = gatemen

    # 11. Keymen
    if 'KEYMAN' in wb.sheetnames:
        ws = wb['KEYMAN']
        keymen = []
        rows = list(ws.iter_rows(values_only=True))
        for r in rows[2:]:
            cells = [clean_val(c) for c in r]
            if len(cells) >= 5 and (cells[1] or cells[2]):
                keymen.append({
                    'code': str(cells[0]),
                    'name': str(cells[2] or cells[1]),
                    'father_name': str(cells[3]),
                    'route': str(cells[4]),
                    'mobile': str(cells[5]) if len(cells) > 5 else '',
                    'rest_giver': str(cells[6]) if len(cells) > 6 else ''
                })
        results['keymen'] = keymen

    # 12. Bridge Watchmen
    if 'Br.Watchman' in wb.sheetnames:
        ws = wb['Br.Watchman']
        watchmen = []
        rows = list(ws.iter_rows(values_only=True))
        for r in rows[3:]:
            cells = [clean_val(c) for c in r]
            if len(cells) >= 3 and cells[1]:
                name_clean = str(cells[1]).split('\n')[0].strip()
                watchmen.append({
                    's_no': cells[0],
                    'name': name_clean,
                    'id_no': str(cells[2]),
                    'mobile': str(cells[3]) if len(cells) > 3 else ''
                })
        results['bridge_watchmen'] = watchmen

    # 13. Patrolmen & Shifts (SPN & SPD)
    if 'SPN & SPD' in wb.sheetnames:
        ws = wb['SPN & SPD']
        patrolmen = []
        rows = list(ws.iter_rows(values_only=True))
        for r in rows[2:]:
            cells = [clean_val(c) for c in r]
            if len(cells) >= 5 and (cells[0] or cells[1]):
                patrolmen.append({
                    'code': str(cells[0]),
                    'name': str(cells[1]),
                    'route': str(cells[3]),
                    'mobile': str(cells[4]),
                    'rest_giver': str(cells[5]) if len(cells) > 5 else ''
                })
        results['patrolmen'] = patrolmen

    return results

def parse_loop_line_data():
    xls_path = BASE_DIR / 'Loop Line Data.xls'
    if not xls_path.exists():
        return []
    
    print(f"Loading {xls_path}...")
    wb = xlrd.open_workbook(str(xls_path))
    sheet = wb.sheet_by_name('Data of each line')
    lines = []
    for r in range(1, sheet.nrows):
        stn = str(sheet.cell_value(r, 0)).strip()
        line = str(sheet.cell_value(r, 1)).strip()
        xover = str(sheet.cell_value(r, 2)).strip()
        km_from = safe_float(sheet.cell_value(r, 3))
        met_from = safe_float(sheet.cell_value(r, 4))
        km_to = safe_float(sheet.cell_value(r, 5))
        met_to = safe_float(sheet.cell_value(r, 6))
        length = safe_float(sheet.cell_value(r, 7))

        ch_from = round(km_from + (met_from / 1000.0), 4) if km_from else 0.0
        ch_to = round(km_to + (met_to / 1000.0), 4) if km_to else 0.0

        if stn or line:
            lines.append({
                'station': stn,
                'line': line,
                'xover': xover,
                'km_from': km_from,
                'met_from': met_from,
                'ch_from': ch_from,
                'km_to': km_to,
                'met_to': met_to,
                'ch_to': ch_to,
                'length_km': length
            })
    return lines

def parse_vivek_profile():
    vivek_path = BASE_DIR / 'Vivek.xlsx'
    if not vivek_path.exists():
        print("Vivek.xlsx not found.")
        return {'points': [], 'gradients': [], 'high_banks': [], 'link_line': []}
    
    print(f"Loading Profile Data from {vivek_path}...")
    wb = openpyxl.load_workbook(vivek_path, data_only=True, read_only=True)
    
    points = []
    if 'Profile Data' in wb.sheetnames:
        ws = wb['Profile Data']
        rows = list(ws.iter_rows(values_only=True))[2:]
        for r in rows:
            if len(r) >= 3 and r[0] is not None and r[1] is not None and r[2] is not None:
                try:
                    km = float(r[0])
                    gl = float(r[1])
                    pfl = float(r[2])
                    h = float(r[3]) if r[3] is not None else pfl - gl
                    source = str(r[5]) if len(r) > 5 and r[5] else ''
                    points.append({
                        'km': round(km, 4),
                        'gl': round(gl, 3),
                        'pfl': round(pfl, 3),
                        'bank': round(h, 3),
                        'source': source
                    })
                except Exception:
                    pass

    # Sort points by chainage
    points.sort(key=lambda p: p['km'])

    # Segment into continuous gradients
    gradients = []
    if points:
        curr_seg = {
            'km_from': points[0]['km'],
            'pfl_from': points[0]['pfl'],
            'km_to': points[0]['km'],
            'pfl_to': points[0]['pfl'],
            'grade': 0.0,
            'points_count': 1
        }
        
        for i in range(1, len(points)):
            p_prev = points[i-1]
            p_curr = points[i]
            d_dist = (p_curr['km'] - p_prev['km']) * 1000.0
            if d_dist <= 0: continue
            grade = (p_curr['pfl'] - p_prev['pfl']) / d_dist
            rounded_grade = round(grade, 5)
            
            if curr_seg['points_count'] == 1:
                curr_seg['grade'] = rounded_grade
                curr_seg['km_to'] = p_curr['km']
                curr_seg['pfl_to'] = p_curr['pfl']
                curr_seg['points_count'] += 1
            elif abs(rounded_grade - curr_seg['grade']) < 0.00015:
                curr_seg['km_to'] = p_curr['km']
                curr_seg['pfl_to'] = p_curr['pfl']
                curr_seg['points_count'] += 1
            else:
                g = curr_seg['grade']
                if abs(g) < 0.00005:
                    label = 'LEVEL'
                    direction = 'LEVEL'
                else:
                    one_in = round(1.0 / abs(g))
                    direction = 'RISING' if g > 0 else 'FALLING'
                    label = f'1 in {one_in} {direction}'
                
                curr_seg['label'] = label
                curr_seg['direction'] = direction
                curr_seg['length_m'] = round((curr_seg['km_to'] - curr_seg['km_from']) * 1000.0, 1)
                gradients.append(curr_seg)
                
                curr_seg = {
                    'km_from': p_prev['km'],
                    'pfl_from': p_prev['pfl'],
                    'km_to': p_curr['km'],
                    'pfl_to': p_curr['pfl'],
                    'grade': rounded_grade,
                    'points_count': 2
                }
                
        g = curr_seg['grade']
        if abs(g) < 0.00005:
            label = 'LEVEL'
            direction = 'LEVEL'
        else:
            one_in = round(1.0 / abs(g))
            direction = 'RISING' if g > 0 else 'FALLING'
            label = f'1 in {one_in} {direction}'
        curr_seg['label'] = label
        curr_seg['direction'] = direction
        curr_seg['length_m'] = round((curr_seg['km_to'] - curr_seg['km_from']) * 1000.0, 1)
        gradients.append(curr_seg)

    # High Bank records
    high_banks = []
    if 'High Bank >6m' in wb.sheetnames:
        ws = wb['High Bank >6m']
        rows = list(ws.iter_rows(values_only=True))[2:]
        for r in rows:
            if len(r) >= 2 and r[0] is not None and r[1] is not None:
                try:
                    km = float(r[0])
                    h = float(r[1])
                    src = str(r[4]) if len(r) > 4 and r[4] else ''
                    high_banks.append({
                        'km': round(km, 4),
                        'height': round(h, 3),
                        'drawing': src
                    })
                except Exception:
                    pass

    # Link Line data
    link_line = []
    if 'Link Line Data' in wb.sheetnames:
        ws = wb['Link Line Data']
        rows = list(ws.iter_rows(values_only=True))[5:]
        for r in rows:
            if len(r) >= 3 and r[0] is not None and r[1] is not None and r[2] is not None:
                try:
                    km = float(r[0])
                    gl = float(r[1])
                    pfl = float(r[2])
                    h = float(r[3]) if r[3] is not None else pfl - gl
                    link_line.append({
                        'km': round(km, 4),
                        'gl': round(gl, 3),
                        'pfl': round(pfl, 3),
                        'bank': round(h, 3),
                        'section': str(r[5]) if len(r) > 5 and r[5] else 'SMUN-RPJ Link Line'
                    })
                except Exception:
                    pass

    link_line.sort(key=lambda p: p['km'])

    # Segment link line into gradients
    link_gradients = []
    if link_line:
        curr_lseg = {
            'km_from': link_line[0]['km'],
            'pfl_from': link_line[0]['pfl'],
            'km_to': link_line[0]['km'],
            'pfl_to': link_line[0]['pfl'],
            'grade': 0.0,
            'points_count': 1
        }
        for i in range(1, len(link_line)):
            p_prev = link_line[i-1]
            p_curr = link_line[i]
            d_dist = (p_curr['km'] - p_prev['km']) * 1000.0
            if d_dist <= 0: continue
            grade = (p_curr['pfl'] - p_prev['pfl']) / d_dist
            rounded_grade = round(grade, 5)

            if curr_lseg['points_count'] == 1:
                curr_lseg['grade'] = rounded_grade
                curr_lseg['km_to'] = p_curr['km']
                curr_lseg['pfl_to'] = p_curr['pfl']
                curr_lseg['points_count'] += 1
            elif abs(rounded_grade - curr_lseg['grade']) < 0.00015:
                curr_lseg['km_to'] = p_curr['km']
                curr_lseg['pfl_to'] = p_curr['pfl']
                curr_lseg['points_count'] += 1
            else:
                g = curr_lseg['grade']
                if abs(g) < 0.00005:
                    label = 'LEVEL'
                    direction = 'LEVEL'
                else:
                    one_in = round(1.0 / abs(g))
                    direction = 'RISING' if g > 0 else 'FALLING'
                    label = f'1 in {one_in} {direction}'
                curr_lseg['label'] = label
                curr_lseg['direction'] = direction
                curr_lseg['length_m'] = round((curr_lseg['km_to'] - curr_lseg['km_from']) * 1000.0, 1)
                link_gradients.append(curr_lseg)
                curr_lseg = {
                    'km_from': p_prev['km'],
                    'pfl_from': p_prev['pfl'],
                    'km_to': p_curr['km'],
                    'pfl_to': p_curr['pfl'],
                    'grade': rounded_grade,
                    'points_count': 2
                }
        g = curr_lseg['grade']
        if abs(g) < 0.00005:
            label = 'LEVEL'
            direction = 'LEVEL'
        else:
            one_in = round(1.0 / abs(g))
            direction = 'RISING' if g > 0 else 'FALLING'
            label = f'1 in {one_in} {direction}'
        curr_lseg['label'] = label
        curr_lseg['direction'] = direction
        curr_lseg['length_m'] = round((curr_lseg['km_to'] - curr_lseg['km_from']) * 1000.0, 1)
        link_gradients.append(curr_lseg)

    # Official Gradients from Google Sheet
    official_gradients = []
    official_link_gradients = []
    csv_path = DATA_DIR / 'official_gradients.csv'
    if csv_path.exists():
        import csv
        with open(csv_path, 'r', encoding='utf-8') as f:
            crows = list(csv.reader(f))
        for r in crows[5:]:
            if len(r) >= 11 and r[1]:
                try:
                    sec = r[1].strip()
                    line_code = r[2].strip()
                    km_f = float(r[3]) + float(r[4])/1000.0
                    km_t = float(r[5]) + float(r[6])/1000.0
                    length_m = float(r[8]) if r[8] else round((km_t - km_f)*1000.0, 1)
                    direction = r[9].strip()
                    grade_val = float(r[10]) if r[10] else 0.0
                    if grade_val > 0:
                        label = f"1 in {int(grade_val)} {direction}"
                        grade_float = (1.0 / grade_val) if 'RISE' in direction.upper() else (-1.0 / grade_val)
                    else:
                        label = "LEVEL"
                        direction = "LEVEL"
                        grade_float = 0.0

                    item = {
                        'section': sec,
                        'line': line_code,
                        'km_from': round(km_f, 4),
                        'km_to': round(km_t, 4),
                        'length_m': length_m,
                        'direction': direction,
                        'grade_val': grade_val,
                        'grade': grade_float,
                        'label': label
                    }
                    if 'Link' in sec:
                        official_link_gradients.append(item)
                    else:
                        official_gradients.append(item)
                except Exception as e:
                    pass

    print(f"Loaded {len(points)} profile points, {len(gradients)} Vivek-calculated gradients, {len(official_gradients)} official gradients, {len(high_banks)} high banks, {len(link_line)} link line points, {len(link_gradients)} link gradients.")
    return {
        'points': points,
        'gradients': official_gradients if official_gradients else gradients,
        'vivek_gradients': gradients,
        'official_gradients': official_gradients,
        'official_link_gradients': official_link_gradients,
        'high_banks': high_banks,
        'link_line': link_line,
        'link_gradients': official_link_gradients if official_link_gradients else link_gradients,
        'vivek_link_gradients': link_gradients
    }

def parse_store_inventories():
    store_dir = DATA_DIR / 'Store' / 'Store'
    inventories = {
        'pway_material': [],
        'tp_material': [],
        'cp_material': [],
        'uniform': []
    }

    pway_path = store_dir / 'P-way Material IMSD SMUN.xlsx'
    if pway_path.exists():
        wb = openpyxl.load_workbook(pway_path, data_only=True, read_only=True)
        idx_sheet = next((s for s in wb.sheetnames if 'index' in s.lower()), None)
        if idx_sheet:
            ws = wb[idx_sheet]
            for r in list(ws.iter_rows(values_only=True))[3:]:
                cells = [clean_val(c) for c in r]
                if len(cells) >= 7 and cells[1]:
                    inventories['pway_material'].append({
                        's_no': cells[0],
                        'particular': str(cells[1]),
                        'page_no': str(cells[2]),
                        'receipt': safe_float(cells[3]),
                        'transfer': safe_float(cells[4]),
                        'issues': safe_float(cells[5]),
                        'balance': safe_float(cells[6]),
                        'unit': str(cells[7]) if len(cells) > 7 else 'NOS.'
                    })

    tp_path = store_dir / 'T&P Material IMSD SMUN.xlsx'
    if tp_path.exists():
        wb = openpyxl.load_workbook(tp_path, data_only=True, read_only=True)
        idx_sheet = next((s for s in wb.sheetnames if 'index' in s.lower()), None)
        if idx_sheet:
            ws = wb[idx_sheet]
            for r in list(ws.iter_rows(values_only=True))[3:]:
                cells = [clean_val(c) for c in r]
                if len(cells) >= 7 and cells[1]:
                    inventories['tp_material'].append({
                        's_no': cells[0],
                        'particular': str(cells[1]),
                        'page_no': str(cells[2]),
                        'receipt': safe_float(cells[3]),
                        'transfer': safe_float(cells[4]),
                        'issues': safe_float(cells[5]),
                        'balance': safe_float(cells[6]),
                        'unit': str(cells[7]) if len(cells) > 7 else 'NOS.'
                    })

    return inventories

def index_staff_media():
    staff_dir = DATA_DIR / 'Store' / 'Staff'
    media_map = {}

    all_images = glob.glob(str(staff_dir / '**' / '*.*'), recursive=True)
    for img_path in all_images:
        ext = os.path.splitext(img_path)[1].lower()
        if ext not in ['.png', '.jpg', '.jpeg', '.webp']:
            continue
        rel_path = os.path.relpath(img_path, BASE_DIR)
        media_url = f"/media/{rel_path}"
        filename = os.path.basename(img_path)

        parent_dir = os.path.basename(os.path.dirname(img_path)).lower()
        is_qr = 'qr' in parent_dir or 'qr' in filename.lower().split()

        num_match = re.findall(r'\b\d{4,6}\b', filename)
        name_clean = re.sub(r'[\d\(\)\._-]', ' ', os.path.splitext(filename)[0]).strip()
        name_clean = re.sub(r'\b(beat|no|rg|shift|photo|qr)\b', '', name_clean, flags=re.IGNORECASE)
        name_clean = ' '.join(name_clean.lower().split())

        for nid in num_match:
            if nid not in media_map:
                media_map[nid] = {}
            media_map[nid]['qr' if is_qr else 'photo'] = media_url

        if name_clean:
            if name_clean not in media_map:
                media_map[name_clean] = {}
            media_map[name_clean]['qr' if is_qr else 'photo'] = media_url

    return media_map

def index_drawings():
    pp_dir = DATA_DIR / 'P&P'
    drawings = []
    all_pdfs = glob.glob(str(pp_dir / '**' / '*.pdf'), recursive=True)

    for pdf_path in all_pdfs:
        rel_path = os.path.relpath(pdf_path, BASE_DIR)
        fname = os.path.basename(pdf_path)
        if fname.startswith('._'):
            continue

        size_kb = round(os.path.getsize(pdf_path) / 1024, 1)

        # Parse DFCC Chainage from filename or parent path
        dfcc_match = re.search(r'DFCC_(?:LINK_)?KM_(\d+\.?\d*)_to_(\d+\.?\d*)', fname)
        if dfcc_match:
            km_from = float(dfcc_match.group(1))
            km_to = float(dfcc_match.group(2))
        else:
            km_match = re.findall(r'KM\s*(\d+\.?\d*)\s*TO\s*KM\s*(\d+\.?\d*)', pdf_path, re.IGNORECASE)
            km_from = float(km_match[0][0]) if km_match else None
            km_to = float(km_match[0][1]) if km_match else None

        sheet_match = re.search(r'Sheet_(\w+)', fname)
        if not sheet_match:
            sheet_match = re.search(r'-(\d{3,4}[A-Za-z]?)\.pdf', fname)
        sheet_no = sheet_match.group(1) if sheet_match else ""

        cat = "Plan & Profile"
        if "LINK" in fname.upper():
            cat = "Link Line Plan & Profile"
        elif "COMBINED" in fname.upper():
            cat = "Combined Plan & Profile"
        elif "LEGEND" in fname.upper():
            cat = "Legend"
        elif "COVER" in fname.upper():
            cat = "Cover Sheet"
        elif "ESP" in pdf_path.upper():
            cat = "Engineering Scale Plan"

        drawings.append({
            'filename': fname,
            'url': f"/media/{rel_path}",
            'size_kb': size_kb,
            'category': cat,
            'sheet_no': sheet_no,
            'km_from': km_from,
            'km_to': km_to,
            'rel_path': rel_path
        })

    esp_dwgs = glob.glob(str(pp_dir / 'ESP' / '*.dwg'))
    stn_esp_meta = {
        'SHAMBHU': {'code': 'SMUN', 'name': 'New Shambhu', 'km_from': 1168.0, 'km_to': 1171.0},
        'BANJARA': {'code': 'SBJN', 'name': 'Sarai Banjara', 'km_from': 1187.0, 'km_to': 1190.0},
        'GOBINDGARH': {'code': 'GVGN', 'name': 'Mandi Gobindgarh', 'km_from': 1211.5, 'km_to': 1215.0},
        'KHANNA': {'code': 'KNNN', 'name': 'New Khanna', 'km_from': 1227.5, 'km_to': 1230.5},
        'CHAWAPAIL': {'code': 'CHAN', 'name': 'New Chawa Pail', 'km_from': 1236.0, 'km_to': 1239.0}
    }

    for dwg in esp_dwgs:
        rel_path = os.path.relpath(dwg, BASE_DIR)
        fname = os.path.basename(dwg)
        fname_up = fname.upper()
        stn_code = None
        stn_name = None
        km_f = None
        km_t = None

        for k, meta in stn_esp_meta.items():
            if k in fname_up:
                stn_code = meta['code']
                stn_name = meta['name']
                km_f = meta['km_from']
                km_t = meta['km_to']
                break

        drawings.append({
            'filename': fname,
            'url': f"/media/{rel_path}",
            'size_kb': round(os.path.getsize(dwg) / 1024, 1),
            'category': "ESP CAD Drawing (AutoCAD)",
            'sheet_no': "ESP",
            'station': stn_code,
            'station_name': stn_name,
            'km_from': km_f,
            'km_to': km_t,
            'rel_path': rel_path
        })

    drawings.sort(key=lambda d: (d.get('km_from') or 9999, d.get('sheet_no') or '', d['filename']))
    return drawings

def parse_trc_summary():
    trc_dir = DATA_DIR / 'Store' / 'TRC' / 'Track_Recording_24.08.26'
    summary = {
        'inspection_date': '24.08.2026',
        'section': 'SNL - UBCD (Single Line)',
        'chainage_range': 'KM 1250 to KM 1159',
        'fastening_defects_count': 106,
        'sleeper_defects_count': 15500,
        'rail_defects_count': 29509,
        'ballast_defects_count': 61161,
        'fastening_sample': []
    }
    
    fastening_file = next((f for f in glob.glob(str(trc_dir / '*Fastening*.xlsx'))), None)
    if fastening_file and os.path.exists(fastening_file):
        try:
            wb = openpyxl.load_workbook(fastening_file, data_only=True, read_only=True)
            ws = wb.active
            rows = list(ws.iter_rows(values_only=True))
            for r in rows[2:25]:
                cells = [clean_val(c) for c in r]
                if any(cells):
                    summary['fastening_sample'].append(cells[:8])
        except Exception as e:
            print("TRC error:", e)
    
    return summary

def main():
    print("Building Vani Unified Railway Database...")
    db = {}

    at_a_glance = parse_at_a_glance()
    db.update(at_a_glance)

    loop_lines = parse_loop_line_data()
    db['loop_lines'] = loop_lines

    profile_data = parse_vivek_profile()
    db['profile'] = profile_data

    store_data = parse_store_inventories()
    db['store'] = store_data

    media_map = index_staff_media()
    db['media_map'] = media_map

    enriched_staff = []
    for member in db.get('master_staff', []):
        m = dict(member)
        awpo = str(m.get('awpo_id', '')).strip()
        name_clean = ' '.join(re.sub(r'[\d\(\)\._-]', ' ', m.get('name', '')).lower().split())

        photo = None
        qr = None

        if awpo and awpo in media_map:
            photo = media_map[awpo].get('photo')
            qr = media_map[awpo].get('qr')

        if not photo or not qr:
            for k, v in media_map.items():
                if not photo and 'photo' in v and (k in name_clean or name_clean in k):
                    photo = v['photo']
                if not qr and 'qr' in v and (k in name_clean or name_clean in k):
                    qr = v['qr']

        m['photo_url'] = photo
        m['qr_url'] = qr
        enriched_staff.append(m)

    db['master_staff'] = enriched_staff

    drawings = index_drawings()
    db['drawings'] = drawings

    trc = parse_trc_summary()
    db['trc'] = trc

    # Load Google Maps KM Coordinates
    km_coords_path = DATA_DIR / 'km_coordinates.json'
    if km_coords_path.exists():
        with open(km_coords_path, 'r', encoding='utf-8') as f:
            db['km_coordinates'] = json.load(f)

    db['summary'] = {
        'total_bridges': len(db.get('bridges', [])),
        'total_curves': len(db.get('curves', [])),
        'total_points': len(db.get('points_and_crossings', [])),
        'total_lwr': len(db.get('lwr_sej', [])),
        'total_loop_lines': len(db.get('loop_lines', [])),
        'total_turnout_defects': len(db.get('turnout_defects', [])),
        'total_dfwo': len(db.get('dfwo', [])),
        'total_master_staff': len(db.get('master_staff', [])),
        'total_permanent_staff': len(db.get('permanent_staff', [])),
        'total_gatemen': len(db.get('gatemen', [])),
        'total_keymen': len(db.get('keymen', [])),
        'total_drawings': len(db.get('drawings', [])),
        'total_profile_points': len(db.get('profile', {}).get('points', [])),
        'total_gradient_segments': len(db.get('profile', {}).get('gradients', [])),
        'total_high_banks': len(db.get('profile', {}).get('high_banks', [])),
        'total_pway_stock_items': len(db.get('store', {}).get('pway_material', [])),
        'total_tp_stock_items': len(db.get('store', {}).get('tp_material', [])),
    }

    with open(OUTPUT_FILE, 'w', encoding='utf-8') as f:
        json.dump(db, f, indent=2, ensure_ascii=False)

    print(f"Database successfully saved to {OUTPUT_FILE}")
    print("Summary:", json.dumps(db['summary'], indent=2))

if __name__ == '__main__':
    main()
